"""Generate a Word (.docx) report from the results produced by
stats_engine.run_analysis."""

import io

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _add_bar_chart(doc: Document, fig_counter: list, titre: str, categories: list, pourcentages: list) -> None:
    """Diagramme en barres numéroté et titré au-dessus (convention CAMES pour
    les illustrations : numérotées, titrées, commentées dans le texte)."""
    fig_counter[0] += 1
    p = doc.add_paragraph()
    p.add_run(f"Figure {fig_counter[0]} — {titre}").bold = True

    fig, ax = plt.subplots(figsize=(5, 3))
    ax.bar([str(c) for c in categories], pourcentages, color="#4472C4")
    ax.set_ylabel("%")
    plt.setp(ax.get_xticklabels(), rotation=30, ha="right")
    fig.tight_layout()

    image_stream = io.BytesIO()
    fig.savefig(image_stream, format="png", dpi=150)
    plt.close(fig)
    image_stream.seek(0)

    doc.add_picture(image_stream, width=Inches(4.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def _add_descriptive_table(doc: Document, tableau_descriptif: dict, fig_counter: list) -> None:
    doc.add_heading("Tableau I — Statistiques descriptives", level=2)

    for var, stats_var in tableau_descriptif.items():
        doc.add_heading(var, level=3)

        if stats_var["type"] == "continue":
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text = "Indicateur", "Valeur"

            rows = [
                ("n", stats_var["n"]),
                ("Moyenne", stats_var["moyenne"]),
                ("Écart-type", stats_var["ecart_type"]),
                ("Médiane", stats_var["mediane"]),
                ("Q1", stats_var["q1"]),
                ("Q3", stats_var["q3"]),
            ]
            for label, value in rows:
                row = table.add_row().cells
                row[0].text, row[1].text = str(label), str(value)
        else:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = "Catégorie", "n", "%"

            for cat, freq in stats_var["frequences"].items():
                row = table.add_row().cells
                row[0].text, row[1].text, row[2].text = str(cat), str(freq["n"]), f"{freq['pct']}%"

            categories = list(stats_var["frequences"].keys())
            pourcentages = [freq["pct"] for freq in stats_var["frequences"].values()]
            _add_bar_chart(doc, fig_counter, f"Répartition de « {var} »", categories, pourcentages)

        doc.add_paragraph()


def _add_test_principal(doc: Document, test: dict) -> None:
    doc.add_heading("Test statistique principal", level=2)

    p = doc.add_paragraph()
    p.add_run("Test utilisé : ").bold = True
    p.add_run(test["test"])

    if test.get("statistic") is not None:
        doc.add_paragraph(f"Statistique : {test['statistic']}")

    doc.add_paragraph(f"p-value : {test['p_value']}")

    p = doc.add_paragraph()
    p.add_run("Résultat : ").bold = True
    p.add_run("différence statistiquement significative (p < 0.05)"
               if test["significatif"]
               else "pas de différence statistiquement significative (p ≥ 0.05)")

    doc.add_paragraph(f"Groupes comparés : {', '.join(test['groupes'])}")

    effect_size = test.get("effect_size")
    if effect_size:
        p = doc.add_paragraph()
        p.add_run("Taille d'effet : ").bold = True
        if "ic95" in effect_size:
            p.add_run(f"{effect_size['nom']} = {effect_size['valeur']} (IC 95% : [{effect_size['ic95'][0]} ; {effect_size['ic95'][1]}])")
        else:
            p.add_run(f"{effect_size['nom']} = {effect_size['valeur']} ({effect_size['interpretation']})")

    if test.get("interpretation"):
        p = doc.add_paragraph()
        p.add_run("En clair : ").bold = True
        p.add_run(test["interpretation"])

    doc.add_paragraph()


def _add_regression(doc: Document, regression: dict) -> None:
    doc.add_heading(f"Analyse multivariée — {regression['type']}", level=2)
    doc.add_paragraph(f"n = {regression['n']}")

    if "r2" in regression:
        doc.add_paragraph(f"R² = {regression['r2']} (R² ajusté = {regression['r2_ajuste']})")
        coef_label = "Coefficient"
    else:
        doc.add_paragraph(f"Pseudo-R² = {regression['pseudo_r2']}")
        doc.add_paragraph(
            f"Catégorie de référence : {regression['categorie_reference']} "
            f"/ événement : {regression['categorie_evenement']}"
        )
        coef_label = "Odds ratio"

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text = (
        "Variable", coef_label, "IC 95%", "p-value", "En clair"
    )

    for coef in regression["coefficients"]:
        row = table.add_row().cells
        value = coef.get("odds_ratio", coef.get("coefficient"))
        row[0].text = coef["variable"]
        row[1].text = str(value)
        row[2].text = f"[{coef['ic95'][0]} ; {coef['ic95'][1]}]"
        row[3].text = str(coef["p_value"])
        row[4].text = coef.get("interpretation", "")

    doc.add_paragraph()


def _add_correlations(doc: Document, correlations: list) -> None:
    doc.add_heading("Corrélations", level=2)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text, hdr[5].text = (
        "Variable", "Méthode", "r", "p-value", "Force", "En clair"
    )

    for corr in correlations:
        row = table.add_row().cells
        if "erreur" in corr:
            row[0].text = corr["variable"]
            row[1].text = corr["erreur"]
            continue
        row[0].text = corr["variable"]
        row[1].text = corr["methode"]
        row[2].text = str(corr["r"])
        row[3].text = str(corr["p_value"])
        row[4].text = corr["force"]
        row[5].text = corr["interpretation"]

    doc.add_paragraph()


def _format_effect_size(effect_size: dict) -> str:
    if not effect_size:
        return ""
    if "ic95" in effect_size:
        return f"{effect_size['nom']} = {effect_size['valeur']} (IC95% [{effect_size['ic95'][0]} ; {effect_size['ic95'][1]}])"
    return f"{effect_size['nom']} = {effect_size['valeur']} ({effect_size.get('interpretation', '')})"


def _add_resume(doc: Document, resume: dict, variable_principale: str) -> None:
    doc.add_heading("Résumé", level=2)

    doc.add_paragraph(
        f"{resume['n_variables_testees']} variable(s) testée(s) en lien avec « {variable_principale} »."
    )

    if resume["variables_significatives"]:
        p = doc.add_paragraph()
        p.add_run(f"{resume['n_associations_significatives']} association(s) significative(s) (p < 0.05) :").bold = True
        for phrase in resume.get("phrases_significatives", []):
            doc.add_paragraph(phrase, style="List Bullet")
    else:
        doc.add_paragraph("Aucune association significative (p < 0.05) détectée.")

    if resume.get("avertissement"):
        p = doc.add_paragraph()
        p.add_run("Important : ").bold = True
        p.add_run(resume["avertissement"])

        if resume["variables_significatives_apres_correction"]:
            p = doc.add_paragraph()
            p.add_run("Après correction de Bonferroni, restent significatives : ").bold = True
            p.add_run(", ".join(resume["variables_significatives_apres_correction"]))
        else:
            doc.add_paragraph("Après correction de Bonferroni, aucune association ne reste significative.")

    doc.add_paragraph()


def _add_associations(doc: Document, associations: list, variable_principale: str) -> None:
    doc.add_heading(f"Tests d'association avec « {variable_principale} »", level=2)

    table = doc.add_table(rows=1, cols=7)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text, hdr[5].text, hdr[6].text = (
        "Variable", "Test", "p-value", "p corrigée", "Significatif", "Taille d'effet", "En clair"
    )

    for assoc in associations:
        row = table.add_row().cells
        row[0].text = assoc["variable"]
        row[1].text = assoc["test"]
        row[2].text = str(assoc["p_value"])
        row[3].text = str(assoc["p_value_corrigee"])
        row[4].text = "Oui" if assoc["significatif"] else "Non"
        row[5].text = _format_effect_size(assoc.get("effect_size"))
        row[6].text = assoc.get("interpretation", "")

    doc.add_paragraph()


def build_auto_report(results: dict) -> io.BytesIO:
    doc = Document()

    variable_principale = results["variable_principale"]
    doc.add_heading("Rapport d'analyse automatique — ThèsIA", level=1)
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph(f"Variable principale étudiée : {variable_principale}")

    fig_counter = [0]

    if "resume" in results:
        _add_resume(doc, results["resume"], variable_principale)

    if results.get("tableau_descriptif"):
        _add_descriptive_table(doc, results["tableau_descriptif"], fig_counter)

    if results.get("associations"):
        _add_associations(doc, results["associations"], variable_principale)

    if "regression" in results:
        _add_regression(doc, results["regression"])
    elif "regression_erreur" in results:
        doc.add_heading("Analyse multivariée", level=2)
        doc.add_paragraph(f"Modèle ajusté non calculé : {results['regression_erreur']}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_report(results: dict) -> io.BytesIO:
    doc = Document()

    doc.add_heading("Rapport d'analyse statistique — ThèsIA", level=1)
    doc.styles["Normal"].font.size = Pt(11)

    fig_counter = [0]

    if "tableau_descriptif" in results and results["tableau_descriptif"]:
        _add_descriptive_table(doc, results["tableau_descriptif"], fig_counter)

    if "test_principal" in results:
        _add_test_principal(doc, results["test_principal"])

    if "regression" in results:
        _add_regression(doc, results["regression"])
    elif "regression_erreur" in results:
        doc.add_heading("Analyse multivariée", level=2)
        doc.add_paragraph(f"Régression non calculée : {results['regression_erreur']}")

    if "correlations" in results:
        _add_correlations(doc, results["correlations"])
    elif "correlations_erreur" in results:
        doc.add_heading("Corrélations", level=2)
        doc.add_paragraph(f"Corrélations non calculées : {results['correlations_erreur']}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
